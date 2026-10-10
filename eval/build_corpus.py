"""Builds the evaluation corpus (1 PDF + 1 PPTX + 2 TXT) from the text below.

The corpus describes a fictional company on purpose: the LLM cannot know these facts
from pre-training, so a correct answer is only possible if retrieval found the right
chunk. It also plants look-alike distractors (several form numbers, several hour
counts, several certification levels) so weak retrieval is visible.

Run:  python eval/build_corpus.py   (needs reportlab + python-pptx: pip install -r requirements-dev.txt)
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).parent / "corpus"

HANDBOOK = [
    ("Halcyon Robotics Field Operations Handbook", """
Version 4.2, effective 1 March 2026. Owner: Field Operations Office (FOO).

Halcyon Robotics builds and operates autonomous inspection robots for utilities, rail operators and
wind-farm owners. This handbook applies to every field operator, crew lead and contractor who deploys
Halcyon equipment at a customer site. Where this handbook and a customer's site rules disagree, the
stricter rule applies.

The company runs two robot families. The HX-7 is a quadcopter inspection drone used for power lines,
towers and turbine blades. The Talon-3 is a tracked ground rover used inside substations, tunnels and
rail yards where flying is not allowed. Both families report telemetry to the Halcyon fleet platform
through regional telemetry gateways.

Operators are organized into regional crews. Each crew has a crew lead who signs off the daily
deployment plan, and a duty manager who can be reached around the clock through the on-call rotation
described in section 8.
"""),
    ("1. Equipment overview: HX-7 and Talon-3", """
HX-7 inspection drone. Maximum takeoff weight 6.8 kg. The HX-7 can stay airborne for up to 38 minutes on
a single battery pack in calm conditions; crews must plan missions for no more than 30 minutes to keep a
reserve. Maximum payload is 1.2 kg, which covers the standard zoom camera or the thermal camera, but not
both at once. Wind limit for operations is 12 m/s sustained or 15 m/s gusts. The HX-7 uses the LX-200
LiDAR module for obstacle avoidance near conductors.

Talon-3 ground rover. Operating time is about 6 hours per battery set, top speed 1.5 m/s, and it can
climb grades up to 25 degrees. The rover carries the same thermal camera as the HX-7 plus an acoustic
sensor used to detect partial discharge in substation equipment. It is rated IP66 and may operate in
rain, but not in standing water deeper than 10 cm.

Both robots are identified by a fleet ID painted on the chassis (for example HX7-0412 or T3-0091). The
fleet ID must be written on every checklist, log and report.
"""),
    ("2. Pre-deployment checks", """
Before every deployment the operator completes the pre-flight checklist on form OPS-12. OPS-12 covers
battery health, propeller condition, firmware version, GPS lock, compass calibration and the geofence
for the site. A deployment may not start until the crew lead has countersigned OPS-12.

Firmware on every robot must be within one minor version of the current fleet release. Robots that
fall further behind are grounded until updated. Compass calibration is required whenever the robot has
been moved more than 50 km since its last flight.

Operators must also confirm that the customer's site permit is valid for the date and that the
emergency landing zones in the deployment plan are clear.
"""),
    ("3. Battery safety", """
All Halcyon robots use lithium-ion battery packs. Packs are stored between 15 and 25 degrees Celsius in
the fireproof cabinets supplied with each crew vehicle. Never charge a pack that is swollen, punctured
or warmer than 45 degrees Celsius; tag it red and isolate it in the sand bin.

For shipping by road or air, packs must be discharged to 40 percent state of charge or less. Packs above
40 percent may not be shipped under any circumstances, even between Halcyon sites.

A pack is retired after 300 charge cycles or when its capacity falls below 80 percent of the rated
value, whichever comes first. Retired packs go back to the regional depot for recycling; they must not
be thrown away with ordinary waste.
"""),
    ("4. Maintenance schedule", """
Routine maintenance on the HX-7 is due every 120 flight hours and is recorded on form OPS-23, the
maintenance log. Propellers are replaced every 300 flight hours or immediately if any nick or crack is
found during the OPS-12 check. Motors are inspected at 600 flight hours.

The Talon-3 tracks are inspected every 200 operating hours and replaced at 1,500 hours. Gearbox oil is
changed once a year.

Only technicians holding certification level L2 or higher may sign off maintenance on form OPS-23.
Level L1 operators may perform the pre-flight checklist but may not open sealed components.
"""),
    ("5. Incident and near-miss reporting", """
Any incident involving injury, property damage, loss of a robot, or a fly-away must be reported within 24
hours on form OPS-19, the incident report. The crew lead also phones the duty manager immediately for
any injury or any loss of control near people.

Near-misses, such as a robot entering a geofence buffer or a battery warning during flight, are reported
within 72 hours on the same OPS-19 form with the box "near-miss" ticked. Reports are blame-free: the goal
is to fix systems, not to punish individuals.

Telemetry outages that affect more than 50 robots are handled as major incidents by the platform
on-call engineer and get a written postmortem within 10 business days.
"""),
    ("6. Data handling and customer privacy", """
Inspection imagery belongs to the customer. Raw video and photos from customer sites are uploaded to the
customer's project space within 48 hours and deleted from robot storage and crew laptops after upload.

Any footage that unintentionally captures people or neighbouring private property is deleted after 30
days unless the customer requests it for an investigation. Footage may never be posted on social media
or used in marketing without written permission from the customer.

Crew laptops use full-disk encryption. A lost laptop or robot storage card is reported as an incident on
OPS-19.
"""),
    ("7. Training and certification", """
Operators progress through three certification levels.

Level L1 requires 8 hours of classroom training and a written exam. L1 operators may fly the HX-7 only
under direct supervision.

Level L2 requires at least 40 supervised flight hours after L1, a practical exam with a senior examiner,
and a first-aid certificate. L2 operators may fly unsupervised in daylight and sign maintenance logs.

Level L3 requires 120 total flight hours, a night-operations exam and a recommendation from a crew lead.
Only L3 operators may fly at night or beyond visual line of sight where the customer's permit allows it.
Certifications are renewed every 24 months.
"""),
    ("8. On-call rotation, travel and expenses", """
Every region keeps one duty manager on call at all times. Shifts run for one week, Monday 09:00 to the
following Monday 09:00 local time. The on-call phone must be answered within 15 minutes.

Travel is booked through the company travel desk. The meal per diem for field travel is 68 US dollars
per day; hotel costs are reimbursed at actual cost up to the city limit published by Finance. Expense
claims are filed on form FIN-4 within 30 days of the trip.

Crew vehicles must carry the fire extinguisher, the sand bin for damaged batteries and the first-aid
kit at all times.
"""),
]

DECK = [
    ("Project Kestrel: Q3 2026 Review", ["Halcyon Robotics · Automated inspection program", "Prepared for the leadership review, 6 October 2026"],
     "Kestrel is the program that moved our utility inspections from manual drone piloting to scheduled autonomous missions."),
    ("Q3 goals", ["Scale autonomous inspections to 1,400 sites per quarter", "Raise defect detection recall above 90%",
                  "Launch the customer reporting portal", "Keep total program spend under budget"],
     "We hit three of the four goals; the portal slipped to Q4."),
    ("Key results", ["Sites inspected: 1,450 (up 22% quarter over quarter)", "Defect detection recall: 91.3% (target 90%)",
                     "Average mission time: 26 minutes", "Customer NPS: 47 (Q2: 41)"],
     "Recall is measured on the labelled validation set of 4,800 images reviewed by our senior inspectors."),
    ("Model improvements", ["Retrained the defect detector on 31,000 new labelled images", "Added thermal hot-spot classifier for transformers",
                            "False positives per 100 images fell from 6.1 to 3.8"],
     "The thermal classifier runs on the Talon-3 rover's thermal camera as well as on the HX-7."),
    ("Budget", ["Approved FY2026 budget: $3.1M", "Spent through Q3: $2.4M", "Largest cost: field crew hours (46%)",
                "Hardware: 18 new HX-7 units purchased"],
     "We expect to finish the year about $150K under budget if the LiDAR deliveries arrive as planned."),
    ("Risks", ["Supply chain: the LX-200 LiDAR module has a 14-week lead time", "Two crews are below staffing target in the North region",
               "Telemetry reliability after INC-2047"],
     "We placed a buffer order of 25 LX-200 modules in September to cover Q1 2027 builds."),
    ("Incident follow-up: INC-2047", ["August telemetry outage took 312 robots offline", "Certificate renewal is now automated",
                                      "Alerting threshold reduced to 3 minutes"],
     "All high-priority action items from the postmortem are closed except the gateway failover drill."),
    ("Customer reporting portal", ["Status: slipped to Q4", "Beta with three utility customers", "Exports inspection findings as PDF and CSV"],
     "The delay came from the security review requested by our largest customer."),
    ("Next steps for Q4", ["Launch the reporting portal to all customers", "Start a pilot with a new water utility in November",
                           "Hire 6 more field operators for the North region", "Run the gateway failover drill"],
     "The November pilot customer is Brightwater Utilities, covering 220 pumping stations in two counties."),
    ("Team and contacts", ["Program lead: Mara Okafor", "Head of ML: Daniel Reyes", "Field operations: Priya Natarajan",
                           "Questions: kestrel-program channel"],
     "Mara Okafor took over as program lead in July after the reorganization."),
]

POSTMORTEM = """POSTMORTEM: INC-2047 Fleet telemetry outage (us-east region)
Status: Final. Author: Jonas Whitfield, platform on-call engineer. Review date: 2026-08-22.
Severity: SEV-1 (more than 50 robots affected)

SUMMARY

On 2026-08-14 the telemetry gateway tg-east-2 stopped accepting connections from robots in the us-east
region. 312 robots (271 HX-7 drones and 41 Talon-3 rovers) lost their connection to the fleet platform
for 3 hours 42 minutes. Robots that were already flying completed their missions and landed using
on-board autonomy, as designed. No robot was lost, nobody was hurt, and no customer data was lost, but
46 scheduled inspections had to be rescheduled.

ROOT CAUSE

The TLS certificate on gateway tg-east-2 expired at 02:00 UTC. The certificate had been issued manually
in 2025 and was not in the automated renewal system, so nobody received the renewal reminder. Robots
correctly refused to connect to a gateway presenting an expired certificate.

TIMELINE (all times UTC, 2026-08-14)

02:00  Certificate on tg-east-2 expires. Robots start failing their TLS handshake and retrying.
02:10  Connection error rate on tg-east-2 passes 90 percent.
02:25  The "gateway connections dropped" alert fires. The alert only fires after the error rate has
       stayed high for 15 minutes, which delayed detection.
02:31  On-call engineer Jonas Whitfield acknowledges the page.
03:05  Root cause identified after ruling out a network problem.
03:40  A new certificate is requested; issuance is slow because the manual process needs two approvers.
05:20  New certificate installed on tg-east-2. Robots begin reconnecting.
05:42  All 312 robots reconnected. Incident declared resolved.

IMPACT

- 312 robots without telemetry for 3 hours 42 minutes (02:00 to 05:42 UTC).
- 46 inspections rescheduled across 9 customers.
- No data loss: robots buffered telemetry locally and uploaded it after reconnecting.
- Estimated cost: 38,500 US dollars in rescheduled crew time.

WHAT WENT WELL

- On-board autonomy worked: every flying robot landed safely at its planned landing zone.
- Local telemetry buffering meant no inspection data was lost.
- The customer success team notified all affected customers within one hour of detection.

WHAT WENT WRONG

- The certificate was outside the automated renewal system.
- The alert threshold of 15 minutes was too slow for a total outage.
- There is no automatic failover to the second us-east gateway, tg-east-1.
- The manual certificate process needs two approvers, one of whom was asleep.

ACTION ITEMS

AI-1  Move all gateway certificates into automated renewal. Owner: Lena Brandt. Due 2026-08-31. Done.
AI-2  Lower the "gateway connections dropped" alert threshold from 15 minutes to 3 minutes.
      Owner: Jonas Whitfield. Due 2026-08-21. Done.
AI-3  Add automatic failover from tg-east-2 to tg-east-1. Owner: Sam Ortiz. Due 2026-10-15. In progress.
AI-4  Run a quarterly gateway failover drill. Owner: Sam Ortiz. First drill due 2026-11-05.
AI-5  Add certificate expiry (30 days ahead) to the weekly platform health report. Owner: Lena Brandt.
      Due 2026-09-05. Done.

LESSONS LEARNED

Manual exceptions to automated systems are where outages hide. Every certificate, key and secret must
be tracked by the same automation, with no "temporary" manual exceptions. Alerts for total outages
should fire in minutes, not after a long smoothing window.
"""


POSTMORTEM_1983 = """POSTMORTEM: INC-1983 Talon-3 firmware rollout halted (eu-west region)
Status: Final. Author: Aiko Tanaka, fleet software lead. Review date: 2026-05-12.
Severity: SEV-2 (fewer than 50 robots affected)

SUMMARY

On 2026-05-03 firmware release 7.4.0 was rolled out to Talon-3 rovers in the eu-west region. After the
update, 37 rovers reported false "track slip" warnings and stopped mid-mission. The rollout was halted
and the rovers were rolled back to firmware 7.3.2. Downtime per rover averaged 2 hours 10 minutes.
No equipment was damaged and nobody was hurt.

ROOT CAUSE

Release 7.4.0 changed the wheel-encoder sampling rate from 100 Hz to 250 Hz, but the slip-detection
threshold was not rescaled. On wet concrete the higher sampling rate produced short spikes that crossed
the old threshold. The staging test site was dry, so the problem was not caught before release.

TIMELINE (all times UTC, 2026-05-03)

08:00  Staged rollout of 7.4.0 begins (25 percent of eu-west Talon-3 fleet).
08:40  First false "track slip" stop reported by a crew in Rotterdam.
09:15  Error-budget alert fires after 10 rovers stop within one hour.
09:20  Rollout paused by on-call engineer Aiko Tanaka.
10:05  Rollback to 7.3.2 starts.
11:30  All 37 affected rovers back on 7.3.2 and resuming missions.

IMPACT

- 37 Talon-3 rovers stopped mid-mission; average downtime 2 hours 10 minutes.
- 12 substation inspections rescheduled across 4 customers.
- Estimated cost: 9,200 US dollars in rescheduled crew time.

ACTION ITEMS

AI-1  Add a wet-surface test track to the staging site. Owner: Aiko Tanaka. Due 2026-06-15. Done.
AI-2  Make slip thresholds scale automatically with sampling rate. Owner: Felix Braun. Due 2026-05-31. Done.
AI-3  Limit the first rollout wave to 10 percent of a region. Owner: Aiko Tanaka. Due 2026-05-20. Done.

LESSONS LEARNED

Test environments must include the conditions robots actually meet in the field, including rain and wet
floors. Staged rollouts worked: only a quarter of the fleet received the faulty release.
"""


def build_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter, title="Halcyon Robotics Field Operations Handbook",
                            author="Halcyon Robotics (fictional)")
    story = []
    for i, (heading, body) in enumerate(HANDBOOK):
        story.append(Paragraph(heading, styles["Title"] if i == 0 else styles["Heading1"]))
        story.append(Spacer(1, 8))
        for para in [p.strip() for p in body.strip().split("\n\n") if p.strip()]:
            story.append(Paragraph(" ".join(para.split()), styles["BodyText"]))
            story.append(Spacer(1, 6))
        if i < len(HANDBOOK) - 1:
            story.append(PageBreak())
    doc.build(story)


def build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    for i, (title, bullets, notes) in enumerate(DECK):
        layout = prs.slide_layouts[0] if i == 0 else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for bullet in bullets[1:]:
            p = body.add_paragraph()
            p.text = bullet
        for p in body.paragraphs:
            for r in p.runs:
                r.font.size = Pt(20)
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf(OUT / "halcyon_field_operations_handbook.pdf")
    build_pptx(OUT / "project_kestrel_q3_review.pptx")
    (OUT / "postmortem_INC-2047.txt").write_text(POSTMORTEM, encoding="utf-8")
    (OUT / "postmortem_INC-1983.txt").write_text(POSTMORTEM_1983, encoding="utf-8")
    print("Corpus written to", OUT)


if __name__ == "__main__":
    main()
